import csv
import gc
import logging
import os
import shutil
import subprocess
import tempfile
import time
from dataclasses import dataclass
from pathlib import Path

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

from backend.services.libreoffice import (
    is_valid_pdf_structure,
    purge_dir,
    reap_process,
    resolve_libreoffice_binary,
    run_libreoffice_subprocess,
)

logger = logging.getLogger("xlsx2pdf")

_ZIP_MAGIC = b"PK\x03\x04"  # .xlsx is a ZIP
_OLE2_MAGIC = b"\xd0\xcf\x11\xe0"  # legacy .xls (unsupported by openpyxl)


class ConversionError(Exception):
    pass


class ConversionTimeout(ConversionError):
    pass


@dataclass
class ConverterConfig:
    timeout_s: float = 45.0
    kill_grace_s: float = 5.0
    soffice_bin: str | None = None
    max_input_bytes: int = 300 * 1024 * 1024
    landscape_col_threshold: int = 7  # >this many cols -> landscape
    lang: str = "C.UTF-8"
    csv_default_encoding: str = "utf-8-sig"


class ExcelToPDFConverter:
    def __init__(self, config: ConverterConfig | None = None, **kw):
        self.cfg = config or ConverterConfig(**kw)
        # Fall back to DocScalar backend's detected LibreOffice path before scanning system PATH
        from backend import config as backend_config

        self.soffice = self.cfg.soffice_bin or getattr(
            backend_config, "LIBREOFFICE_PATH", None
        )

        # Verify binary resolution exists
        resolved = None
        if self.soffice:
            resolved = shutil.which(self.soffice) or (
                Path(self.soffice).exists() and self.soffice
            )
        if not resolved:
            try:
                self.soffice = resolve_libreoffice_binary()
            except Exception as e:
                raise ConversionError(
                    "LibreOffice binary not found. Please install LibreOffice locally "
                    "and ensure it is on your system PATH to run Office conversions."
                ) from e
        else:
            self.soffice = resolved
        logger.info("Using LibreOffice binary: %s", self.soffice)

    # ---------- public entrypoint ----------
    def convert(self, input_path: str, output_path: str) -> str:
        start = time.monotonic()
        in_path = Path(input_path).resolve()
        out_path = Path(output_path).resolve()
        self._preflight(in_path)

        tmpdir = tempfile.mkdtemp(prefix="xlsx2pdf_")
        proc: subprocess.Popen | None = None
        stderr_data = b""
        try:
            # 1/2/3: build a LibreOffice-friendly working copy
            ext = in_path.suffix.lower()
            if ext == ".csv":
                work_in = self._prepare_csv(in_path, tmpdir)
            else:
                work_in = self._prepare_xlsx(in_path, tmpdir)

            profile_dir = Path(tmpdir) / "lo_profile"
            profile_dir.mkdir(parents=True, exist_ok=True)
            out_dir = Path(tmpdir) / "out"
            out_dir.mkdir(parents=True, exist_ok=True)

            cmd = self._build_command(work_in, profile_dir, out_dir)
            env = self._build_env(profile_dir)

            logger.info(
                "Converting %s (timeout=%.0fs)", in_path.name, self.cfg.timeout_s
            )
            proc, _stdout, stderr_data = self._run(cmd, env, tmpdir)

            produced = out_dir / (Path(work_in).stem + ".pdf")
            self._postflight(produced, proc.returncode, stderr_data)

            out_path.parent.mkdir(parents=True, exist_ok=True)
            shutil.move(str(produced), str(out_path))

            dur = time.monotonic() - start
            logger.info(
                "OK %s -> %s in %.2fs (%d bytes)",
                in_path.name,
                out_path.name,
                dur,
                out_path.stat().st_size,
            )
            return str(out_path)

        except ConversionError:
            raise
        except Exception as e:
            logger.exception("Unexpected failure for %s", in_path.name)
            raise ConversionError(str(e)) from e
        finally:
            self._reap(proc)
            self._purge(tmpdir)
            gc.collect()
            logger.debug("Cleanup done; total %.2fs", time.monotonic() - start)

    # ---------- pre-flight validation ----------
    def _preflight(self, in_path: Path):
        if not in_path.exists() or not in_path.is_file():
            raise ConversionError(f"Input not found / not a file: {in_path}")
        size = in_path.stat().st_size
        if size == 0:
            raise ConversionError(f"Input is empty: {in_path}")
        if size > self.cfg.max_input_bytes:
            raise ConversionError(f"Input too large: {size} bytes")
        ext = in_path.suffix.lower()
        if ext == ".csv":
            return
        with open(in_path, "rb") as f:
            head = f.read(8)
        if head.startswith(_OLE2_MAGIC):
            raise ConversionError(
                "Legacy .xls (OLE2) not supported by openpyxl pre-flight. "
                "Convert to .xlsx first or route to a raw-LibreOffice path."
            )
        if not head.startswith(_ZIP_MAGIC):
            raise ConversionError(f"Not a valid .xlsx (bad magic bytes): {in_path}")

    # ---------- 1/2/3: openpyxl pre-processing ----------
    def _prepare_xlsx(self, in_path: Path, tmpdir: str) -> Path:
        try:
            wb = openpyxl.load_workbook(in_path, data_only=False, keep_vba=False)
        except Exception as e:
            raise ConversionError(f"openpyxl could not open workbook: {e}") from e

        try:
            # 3) drop hidden / veryHidden sheets, but keep >=1 visible
            visible = [ws for ws in wb.worksheets if ws.sheet_state == "visible"]
            if not visible:
                raise ConversionError("Workbook has no visible sheets to render.")
            for ws in list(wb.worksheets):
                if ws.sheet_state != "visible":
                    logger.info("Removing %s sheet '%s'", ws.sheet_state, ws.title)
                    del wb[ws.title]

            # 1) force print setup on every remaining visible sheet
            for ws in wb.worksheets:
                self._apply_print_setup(ws)

            # 2) force recalc on load so LibreOffice re-evaluates formulas
            try:
                wb.calculation.fullCalcOnLoad = True
            except Exception:
                logger.debug("Could not set fullCalcOnLoad; relying on LO flag")

            work_in = Path(tmpdir) / in_path.name
            wb.save(work_in)
            return work_in
        finally:
            try:
                wb.close()
            except Exception:
                pass

    def _apply_print_setup(self, ws: Worksheet):
        # measure the real used width
        max_col = ws.max_column or 1
        # clear conflicts: stale print_area and a fixed scale both defeat fitToWidth
        try:
            ws.print_area = None
        except Exception:
            pass
        ws.page_setup.scale = None  # scale and fitTo are mutually exclusive
        ws.sheet_properties.pageSetUpPr.fitToPage = True  # honor fit settings
        ws.page_setup.fitToWidth = 1  # squeeze all columns to 1 page wide
        ws.page_setup.fitToHeight = 0  # allow as many pages tall as needed
        if max_col > self.cfg.landscape_col_threshold:
            ws.page_setup.orientation = "landscape"
        else:
            ws.page_setup.orientation = "portrait"
        # repeat header row on every page for readability
        if ws.max_row and ws.max_row > 1:
            ws.print_title_rows = "1:1"

        # Force grid lines to show in both edit view and printed/exported PDF
        try:
            if ws.views.sheetView:
                ws.views.sheetView[0].showGridLines = True
            else:
                ws.sheet_view.showGridLines = True
        except Exception:
            try:
                ws.sheet_view.showGridLines = True
            except Exception:
                pass
        ws.print_options.gridLines = True

    def _prepare_csv(self, in_path: Path, tmpdir: str) -> Path:
        # sniff delimiter so LibreOffice doesn't mis-split; re-emit clean UTF-8 CSV
        enc = self.cfg.csv_default_encoding
        try:
            with open(in_path, "r", encoding=enc, errors="replace", newline="") as f:
                sample = f.read(8192)
                try:
                    dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
                except csv.Error:
                    dialect = csv.excel
                f.seek(0)
                rows = list(csv.reader(f, dialect))
        except Exception as e:
            raise ConversionError(f"Could not parse CSV: {e}") from e

        work_in = Path(tmpdir) / (in_path.stem + ".csv")
        with open(work_in, "w", encoding="utf-8-sig", newline="") as f:
            csv.writer(f).writerows(rows)
        logger.info("Normalized CSV (%d rows) for conversion", len(rows))
        return work_in

    # ---------- command + environment ----------
    def _build_command(
        self, work_in: Path, profile_dir: Path, out_dir: Path
    ) -> list[str]:
        profile_uri = profile_dir.resolve().as_uri()
        # impress? no -> calc PDF export filter
        return [
            self.soffice,
            "--headless",
            "--nologo",
            "--nofirststartwizard",
            "--nodefault",
            "--norestore",
            "--nolockcheck",
            "--calc",
            f"-env:UserInstallation={profile_uri}",
            "--convert-to",
            "pdf:calc_pdf_Export",
            "--outdir",
            str(out_dir),
            str(work_in),
        ]

    def _build_env(self, profile_dir: Path) -> dict:
        env = os.environ.copy()
        env["HOME"] = str(profile_dir)
        env["LC_ALL"] = self.cfg.lang
        env["LANG"] = self.cfg.lang
        env.pop("DISPLAY", None)
        return env

    # ---------- subprocess with strict timeout + process-group kill ----------
    def _run(self, cmd: list[str], env: dict, cwd: str):
        try:
            return run_libreoffice_subprocess(
                cmd, env, cwd, self.cfg.timeout_s, self.cfg.kill_grace_s
            )
        except subprocess.TimeoutExpired as e:
            raise ConversionTimeout(
                f"LibreOffice exceeded {self.cfg.timeout_s}s"
            ) from e

    # ---------- post-flight ----------
    def _postflight(self, produced: Path, returncode, stderr: bytes):
        err = (stderr or b"").decode("utf-8", "replace").strip()
        if not produced.exists():
            raise ConversionError(
                f"No output produced (exit={returncode}). stderr: {err[:2000]}"
            )
        if produced.stat().st_size == 0:
            raise ConversionError(
                f"Output PDF empty (exit={returncode}). stderr: {err[:2000]}"
            )
        if not self._is_valid_pdf(produced):
            raise ConversionError(
                f"Output not a valid PDF (exit={returncode}). stderr: {err[:2000]}"
            )
        if returncode not in (0, None):
            logger.warning(
                "Non-zero exit %s but valid PDF; stderr: %s", returncode, err[:500]
            )

    @staticmethod
    def _is_valid_pdf(path: Path) -> bool:
        return is_valid_pdf_structure(path)

    # ---------- cleanup ----------
    @staticmethod
    def _reap(proc: subprocess.Popen | None):
        reap_process(proc)

    @staticmethod
    def _purge(tmpdir: str):
        purge_dir(tmpdir)
