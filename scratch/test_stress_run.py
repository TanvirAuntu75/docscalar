import os
import sys
from pathlib import Path
import openpyxl

# Add root folder to sys.path
sys.path.append(str(Path(__file__).resolve().parent.parent))

from backend.services.pdf_to_excel import convert_pdf_to_excel, PDFToExcelConfig, SmartPDFToExcelConverter

def run_stress_test():
    print("==================================================")
    print("DocForge PDF-to-Excel SLANet Table Recognition & Cell Merge Test")
    print("==================================================")
    
    # 1. Locate PDF
    pdf_paths = [
        Path("backend/test_output/multi_table_test.pdf"),
        Path("backend/uploads/13e0a4fa-d068-4fd3-897c-1076e708555a/stress_test_scanned_ocr__1_.pdf"),
        Path("test_dummy.pdf")
    ]
    
    selected_pdf = None
    for p in pdf_paths:
        if p.exists():
            selected_pdf = p
            break
            
    if not selected_pdf:
        print("Error: No test PDF files found.")
        return
        
    print(f"Using PDF: {selected_pdf} ({selected_pdf.stat().st_size} bytes)")
    
    # 2. Run conversion
    print("Converting PDF to Excel (with layout detection enabled)...")
    config = PDFToExcelConfig(enable_layout_detection=True, device="cpu")
    converter = SmartPDFToExcelConverter(config=config)
    
    t0 = openpyxl.Workbook() # just dummy import test
    import time
    start_time = time.time()
    
    result = converter.convert(str(selected_pdf))
    duration = time.time() - start_time
    
    print(f"Conversion result: success={result.success}")
    if not result.success:
        print(f"Error: {result.error_message}")
        return
        
    out_path = Path(result.output_path)
    print(f"Output Excel file created at: {out_path} ({out_path.stat().st_size} bytes)")
    print(f"Duration: {duration:.2f} seconds")
    
    # 3. Verify merges using openpyxl
    print("\nVerifying cell merges in the output Excel sheet...")
    wb = openpyxl.load_workbook(out_path)
    print(f"Sheets in workbook: {wb.sheetnames}")
    
    for name in wb.sheetnames:
        ws = wb[name]
        merged_ranges = list(ws.merged_cells.ranges)
        print(f"Sheet '{name}' has {len(merged_ranges)} merged cell ranges:")
        for r in merged_ranges:
            print(f"  - {r}")
            
    print("==================================================")
    print("Test run complete.")
    print("==================================================")

if __name__ == "__main__":
    run_stress_test()
